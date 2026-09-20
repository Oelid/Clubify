"""Génère le classeur de suivi des tests de Clubify.

Le classeur ne se saisit pas : il se régénère. Ce qui s'y trouve vient de trois
endroits, et de nulle part ailleurs :

- les **scénarios**, extraits du code (`@DisplayName` au backend, `it(...)` au
  frontend) : un test renommé change le classeur, un test supprimé en disparaît ;
- les **résultats**, lus dans les rapports JUnit d'une exécution réelle : aucun
  état n'est déclaré à la main ;
- ce que le code ne peut pas dire — tests à compléter, recette manuelle,
  anomalies — tenu dans `docs/suivi-tests.md`.

Usage :

    python tools/generer-suivi-tests.py              # lit les rapports existants
    python tools/generer-suivi-tests.py --executer   # relance les suites d'abord

Sans rapport disponible, les scénarios sont marqués « Non exécuté » plutôt que
supposés verts : un classeur qui ment sur un test est pire qu'un classeur vide.

Dépendance : openpyxl.
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).resolve().parent.parent
SOURCE = RACINE / "docs" / "suivi-tests.md"
FEATURES = RACINE / "docs" / "features"
RAPPORTS = RACINE / "reports"
SORTIE = RACINE / "docs" / "suivi-tests.xlsx"

ETATS = ["Vert", "Rouge", "Ignoré", "Non exécuté", "À écrire", "Manuel"]

COULEURS_ETAT = {
    "Vert": "C6E7C6",
    "Rouge": "F5C6C6",
    "Ignoré": "E8E8E8",
    "Non exécuté": "FDEBC8",
    "À écrire": "FDEBC8",
    "Manuel": "D8E4F0",
}

TITRE = Font(bold=True, color="FFFFFF")
FOND_TITRE = PatternFill("solid", fgColor="2E4057")
GRAS = Font(bold=True)
BORDURE = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
HAUT = Alignment(vertical="top", wrap_text=True)


@dataclass
class Scenario:
    """Un scénario de test, tel que le code le nomme."""

    feature: str
    reference: str          # critère d'acceptation, quand le nom en cite un
    intitule: str
    niveau: str
    suite: str
    fichier: str
    cle: str = ""                # nom sous lequel le rapport JUnit désigne le cas
    etat: str = "Non exécuté"
    duree: float | None = None
    detail: str = ""


@dataclass
class Feature:
    code: str
    nom: str = ""
    criteres: list[str] = field(default_factory=list)


# --------------------------------------------------------------- extraction

DISPLAY_NAME = re.compile(r'@DisplayName\(\s*"((?:[^"\\]|\\.)*)"\s*\)')
CLASSE = re.compile(r"^\s*(?:public\s+)?(?:final\s+)?class\s+(\w+)", re.M)
# Surefire nomme un cas par sa méthode : le libellé seul ne suffit pas à l'apparier.
NOMME_ET_METHODE = re.compile(
    r'@DisplayName\(\s*"((?:[^"\\]|\\.)*)"\s*\)'
    r'(?:\s*@\w+(?:\([^)]*\))?)*\s*'
    r'(?:public\s+|private\s+|protected\s+)?void\s+(\w+)\s*\(')
IT = re.compile(r"^\s*(?:it|test)(?:\.\w+)?\(\s*(['\"`])((?:(?!\1).)*)\1", re.M)
DESCRIBE = re.compile(r"^\s*(?:describe|test\.describe)\(\s*(['\"`])((?:(?!\1).)*)\1", re.M)
REFERENCE = re.compile(r"^(C\d+[a-z]?)\s*[—-]\s*(.*)$")


def deschapper(texte: str) -> str:
    """Rend lisibles les échappements du code source."""
    return (texte.replace("\\u2019", "’").replace("\\'", "'")
            .replace('\\"', '"').replace("\\n", " "))


def reference_et_intitule(brut: str) -> tuple[str, str]:
    trouve = REFERENCE.match(brut)
    return (trouve.group(1), trouve.group(2)) if trouve else ("", brut)


def scenarios_backend(rattachement) -> list[Scenario]:
    racine = RACINE / "backend" / "src" / "test" / "java"
    trouves: list[Scenario] = []

    for fichier in sorted(racine.rglob("*.java")):
        texte = fichier.read_text(encoding="utf-8")
        if "@Test" not in texte:
            continue

        relatif = fichier.relative_to(RACINE).as_posix()
        classe = CLASSE.search(texte)
        suite = classe.group(1) if classe else fichier.stem

        # Un @DisplayName posé avant la classe nomme la suite ; les autres, un test.
        position_classe = classe.start() if classe else 0
        avant_classe = [n for n in DISPLAY_NAME.finditer(texte) if n.start() < position_classe]
        if avant_classe:
            suite = deschapper(avant_classe[-1].group(1))
        classe_simple = classe.group(1) if classe else fichier.stem

        niveau = ("Architecture" if "ArchitectureTest" in classe_simple
                  else "Intégration (API + base réelle)")

        for brut, methode in NOMME_ET_METHODE.findall(texte):
            reference, intitule = reference_et_intitule(deschapper(brut))
            trouves.append(Scenario(
                feature=rattachement(relatif),
                reference=reference,
                intitule=intitule,
                niveau=niveau,
                suite=suite,
                fichier=relatif,
                cle=f"{classe_simple}#{methode}",
            ))
    return trouves


def scenarios_frontend(rattachement) -> list[Scenario]:
    trouves: list[Scenario] = []
    racines = [RACINE / "frontend" / "projects", RACINE / "frontend" / "e2e"]

    for racine in racines:
        if not racine.exists():
            continue
        for fichier in sorted(racine.rglob("*.spec.ts")):
            if "node_modules" in fichier.parts or "api-client" in fichier.parts:
                continue
            texte = fichier.read_text(encoding="utf-8")
            relatif = fichier.relative_to(RACINE).as_posix()

            bloc = DESCRIBE.search(texte)
            suite = deschapper(bloc.group(2)) if bloc else fichier.stem

            if "/e2e/" in relatif:
                niveau = "Bout en bout (Playwright)"
            elif "/projects/ui/" in relatif:
                niveau = "Unitaire (système de design)"
            else:
                niveau = "Unitaire (interface)"

            for _, brut in IT.findall(texte):
                reference, intitule = reference_et_intitule(deschapper(brut))
                trouves.append(Scenario(
                    feature=rattachement(relatif),
                    reference=reference,
                    intitule=intitule,
                    niveau=niveau,
                    suite=suite,
                    fichier=relatif,
                    cle=f"{suite}#{intitule}",
                ))
    return trouves


# ----------------------------------------------------------------- résultats

def executer_les_suites() -> None:
    """Relance les deux suites et écrit leurs rapports JUnit."""
    RAPPORTS.mkdir(exist_ok=True)
    commandes = [
        (RACINE / "backend", ["cmd", "/c", "mvnw.cmd", "-o", "test"]),
        (RACINE / "frontend", ["cmd", "/c", "npx", "ng", "test", "backoffice",
                               "--watch=false", "--reporters", "junit",
                               "--output-file", str(RAPPORTS / "frontend-backoffice.xml")]),
        (RACINE / "frontend", ["cmd", "/c", "npx", "ng", "test", "ui",
                               "--watch=false", "--reporters", "junit",
                               "--output-file", str(RAPPORTS / "frontend-ui.xml")]),
    ]
    for dossier, commande in commandes:
        print("  exécution :", " ".join(commande[2:4]), "dans", dossier.name)
        subprocess.run(commande, cwd=dossier, check=False,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def lire_les_rapports() -> tuple[dict[str, tuple[str, float, str]], datetime | None]:
    """Résultats indexés par clé de cas, et date de l'exécution la plus récente."""
    resultats: dict[str, tuple[str, float, str]] = {}
    dernier: datetime | None = None

    fichiers = list((RACINE / "backend" / "target" / "surefire-reports").glob("TEST-*.xml"))
    fichiers += list(RAPPORTS.glob("*.xml")) if RAPPORTS.exists() else []

    for fichier in fichiers:
        horodatage = datetime.fromtimestamp(fichier.stat().st_mtime)
        dernier = horodatage if dernier is None else max(dernier, horodatage)
        try:
            racine = ET.parse(fichier).getroot()
        except ET.ParseError:
            continue

        for suite in racine.iter("testsuite"):
            # Surefire nomme la suite par la classe Java ; l'attribut « classname »
            # d'un cas, lui, porte le @DisplayName, qui ne suffit pas à l'identifier.
            classe = (suite.get("name") or "").split(".")[-1]

            for cas in suite.iter("testcase"):
                nom = deschapper(cas.get("name") or "")

                if " > " in nom:
                    # Vitest : « Bloc > intitulé ». Surefire : le nom de la méthode.
                    bloc, _, feuille = nom.rpartition(" > ")
                    _, intitule = reference_et_intitule(feuille.strip())
                    cle = f"{bloc.split(' > ')[0].strip()}#{intitule}"
                else:
                    cle = f"{classe}#{nom.strip()}"

                if cas.find("skipped") is not None:
                    etat, detail = "Ignoré", ""
                elif cas.find("failure") is not None or cas.find("error") is not None:
                    echec = cas.find("failure") if cas.find("failure") is not None else cas.find("error")
                    etat, detail = "Rouge", (echec.get("message") or "")[:300]
                else:
                    etat, detail = "Vert", ""

                resultats[cle] = (etat, float(cas.get("time") or 0), detail)

    return resultats, dernier


# -------------------------------------------------------------------- source

def tableaux_markdown(texte: str) -> dict[str, list[dict[str, str]]]:
    """Les tableaux du fichier source, indexés par le titre qui les précède."""
    tableaux: dict[str, list[dict[str, str]]] = {}
    titre = ""
    entetes: list[str] | None = None

    for ligne in texte.splitlines():
        if ligne.startswith("## "):
            titre, entetes = ligne[3:].strip(), None
            continue
        if not ligne.strip().startswith("|"):
            entetes = None
            continue

        cellules = [c.strip() for c in ligne.strip().strip("|").split("|")]
        if entetes is None:
            entetes = cellules
            tableaux.setdefault(titre, [])
            continue
        if all(set(c) <= {"-", ":", " "} for c in cellules):
            continue
        tableaux[titre].append(dict(zip(entetes, cellules)))

    return tableaux


def criteres_de_la_fiche(chemin: Path) -> tuple[str, list[str]]:
    """Le nom de la feature et ses critères d'acceptation, dans l'ordre."""
    texte = chemin.read_text(encoding="utf-8")
    nom = texte.splitlines()[0].lstrip("# ").strip()
    nom = re.sub(r"^Feature\s*:\s*", "", nom)
    nom = re.sub(r"^F\d+\s*[—-]\s*", "", nom)
    dans_les_criteres = False
    criteres: list[str] = []

    for ligne in texte.splitlines():
        if ligne.startswith("## "):
            dans_les_criteres = ligne.startswith("## Critères d'acceptation")
            continue
        if not dans_les_criteres:
            continue
        for reference in re.findall(r"\b(C\d+[a-z]?)\b", ligne):
            if reference not in criteres:
                criteres.append(reference)

    return nom, criteres


# -------------------------------------------------------------- mise en page

def entete(feuille, colonnes: list[str], largeurs: list[int], ligne: int = 1) -> None:
    for index, (nom, largeur) in enumerate(zip(colonnes, largeurs), start=1):
        cellule = feuille.cell(row=ligne, column=index, value=nom)
        cellule.font = TITRE
        cellule.fill = FOND_TITRE
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
        feuille.column_dimensions[get_column_letter(index)].width = largeur
    feuille.freeze_panes = feuille.cell(row=ligne + 1, column=1)


def colorer_les_etats(feuille, colonne: str, premiere: int, derniere: int) -> None:
    plage = f"{colonne}{premiere}:{colonne}{max(derniere, premiere)}"
    for etat, couleur in COULEURS_ETAT.items():
        feuille.conditional_formatting.add(plage, CellIsRule(
            operator="equal", formula=[f'"{etat}"'],
            fill=PatternFill("solid", fgColor=couleur)))


def border_tout(feuille, premiere: int, derniere: int, colonnes: int) -> None:
    for ligne in range(premiere, derniere + 1):
        for colonne in range(1, colonnes + 1):
            cellule = feuille.cell(row=ligne, column=colonne)
            cellule.border = BORDURE
            cellule.alignment = HAUT


# ------------------------------------------------------------------ feuilles

def feuille_synthese(wb, features, scenarios, a_ecrire, manuels, anomalies, execute) -> None:
    feuille = wb.active
    feuille.title = "Synthèse"

    feuille["A1"] = "Suivi des tests — Clubify"
    feuille["A1"].font = Font(bold=True, size=14)
    feuille["A2"] = ("Régénéré par tools/generer-suivi-tests.py. Les scénarios viennent "
                     "du code, les résultats d'une exécution réelle. Ne pas saisir ici : "
                     "ce qui est manuel vit dans docs/suivi-tests.md.")
    feuille["A2"].alignment = Alignment(wrap_text=True)
    feuille.merge_cells("A2:H2")
    feuille.row_dimensions[2].height = 30
    feuille["A3"] = ("Dernière exécution : "
                     + (execute.strftime("%d/%m/%Y %H:%M") if execute else "aucune"))
    feuille["A3"].font = GRAS

    colonnes = ["Feature", "Nom", "Scénarios", "Verts", "Rouges", "À écrire",
                "Recette manuelle", "Critères couverts", "Anomalies ouvertes"]
    entete(feuille, colonnes, [10, 34, 12, 9, 9, 10, 18, 18, 18], ligne=5)

    ligne = 6
    for feature in features.values():
        siens = [s for s in scenarios if s.feature == feature.code]
        couverts = {s.reference for s in siens if s.reference}
        ouvertes = [a for a in anomalies
                    if a.get("Feature") == feature.code and a.get("État") != "Close"]

        valeurs = [
            feature.code,
            feature.nom,
            len(siens),
            sum(1 for s in siens if s.etat == "Vert"),
            sum(1 for s in siens if s.etat == "Rouge"),
            sum(1 for a in a_ecrire if a.get("Feature") == feature.code),
            sum(1 for m in manuels if m.get("Feature") == feature.code),
            (f"{len(couverts & set(feature.criteres))} / {len(feature.criteres)}"
             if feature.criteres else "—"),
            len(ouvertes),
        ]
        for index, valeur in enumerate(valeurs, start=1):
            feuille.cell(row=ligne, column=index, value=valeur)
        ligne += 1

    border_tout(feuille, 5, ligne - 1, len(colonnes))

    feuille.cell(row=ligne + 1, column=1, value="Niveaux de test").font = GRAS
    par_niveau = defaultdict(int)
    for scenario in scenarios:
        par_niveau[scenario.niveau] += 1
    for decalage, (niveau, nombre) in enumerate(sorted(par_niveau.items()), start=2):
        feuille.cell(row=ligne + decalage, column=1, value=niveau)
        feuille.cell(row=ligne + decalage, column=3, value=nombre)

    depart = ligne + len(par_niveau) + 3
    feuille.cell(row=depart, column=1, value="Lecture des états").font = GRAS
    explications = {
        "Vert": "Passé à la dernière exécution",
        "Rouge": "Échoué à la dernière exécution",
        "Ignoré": "Désactivé ; à justifier",
        "Non exécuté": "Présent dans le code, absent du dernier rapport",
        "À écrire": "Attendu, pas encore écrit (docs/suivi-tests.md)",
        "Manuel": "Déroulé par une personne, pas par la machine",
    }
    for decalage, (etat, texte) in enumerate(explications.items(), start=1):
        cellule = feuille.cell(row=depart + decalage, column=1, value=etat)
        cellule.fill = PatternFill("solid", fgColor=COULEURS_ETAT[etat])
        feuille.cell(row=depart + decalage, column=2, value=texte)


def feuille_feature(wb, feature, scenarios, a_ecrire, manuels) -> None:
    feuille = wb.create_sheet(feature.code)
    feuille["A1"] = f"{feature.code} — {feature.nom}"
    feuille["A1"].font = Font(bold=True, size=13)

    colonnes = ["Réf. critère", "Scénario", "Niveau", "Suite", "État",
                "Durée (s)", "Fichier", "Détail de l'échec"]
    entete(feuille, colonnes, [12, 62, 26, 28, 13, 10, 52, 40], ligne=3)

    ligne = 4
    siens = sorted([s for s in scenarios if s.feature == feature.code],
                   key=lambda s: (s.niveau, s.suite, s.reference or "zz", s.intitule))
    for scenario in siens:
        valeurs = [scenario.reference, scenario.intitule, scenario.niveau, scenario.suite,
                   scenario.etat, round(scenario.duree, 3) if scenario.duree else None,
                   scenario.fichier, scenario.detail]
        for index, valeur in enumerate(valeurs, start=1):
            feuille.cell(row=ligne, column=index, value=valeur)
        ligne += 1

    for source, etat in ((a_ecrire, "À écrire"), (manuels, "Manuel")):
        for entree in [e for e in source if e.get("Feature") == feature.code]:
            intitule = entree.get("Scénario attendu") or entree.get("Scénario", "")
            detail = (entree.get("Pourquoi il manque") or entree.get("Remarque", ""))
            quand = entree.get("Quand") or entree.get("Qui", "")
            valeurs = [None, intitule, entree.get("Niveau", "Recette manuelle"),
                       quand, etat, None, "docs/suivi-tests.md", detail]
            for index, valeur in enumerate(valeurs, start=1):
                feuille.cell(row=ligne, column=index, value=valeur)
            ligne += 1

    colorer_les_etats(feuille, "E", 4, ligne - 1)
    border_tout(feuille, 3, ligne - 1, len(colonnes))


def feuille_couverture(wb, features, scenarios) -> None:
    feuille = wb.create_sheet("Couverture des critères")
    feuille["A1"] = ("Un critère d'acceptation sans test est un critère que personne "
                     "ne vérifie. Les références viennent des fiches de features.")
    feuille["A1"].font = GRAS

    colonnes = ["Feature", "Critère", "Couvert par", "Nb tests", "État"]
    entete(feuille, colonnes, [10, 12, 80, 10, 14], ligne=3)

    ligne = 4
    for feature in features.values():
        par_reference = defaultdict(list)
        for scenario in scenarios:
            if scenario.feature == feature.code and scenario.reference:
                par_reference[scenario.reference].append(scenario)

        for critere in feature.criteres:
            tests = par_reference.get(critere, [])
            etat = "Non exécuté"
            if tests:
                etats = {t.etat for t in tests}
                etat = "Rouge" if "Rouge" in etats else (
                    "Vert" if etats == {"Vert"} else "Non exécuté")
            else:
                etat = "À écrire"

            valeurs = [feature.code, critere,
                       " ; ".join(f"{t.suite} — {t.intitule}" for t in tests) or "aucun test",
                       len(tests), etat]
            for index, valeur in enumerate(valeurs, start=1):
                feuille.cell(row=ligne, column=index, value=valeur)
            ligne += 1

        # Un test qui cite un critère absent de la fiche : l'un des deux a tort.
        for reference in sorted(set(par_reference) - set(feature.criteres)):
            tests = par_reference[reference]
            valeurs = [feature.code, reference,
                       " ; ".join(f"{t.suite} — {t.intitule}" for t in tests),
                       len(tests), "Hors fiche"]
            for index, valeur in enumerate(valeurs, start=1):
                feuille.cell(row=ligne, column=index, value=valeur)
            ligne += 1

    colorer_les_etats(feuille, "E", 4, ligne - 1)
    border_tout(feuille, 3, ligne - 1, len(colonnes))


def feuille_anomalies(wb, anomalies) -> None:
    feuille = wb.create_sheet("Anomalies")
    feuille["A1"] = ("Une anomalie n'est close que lorsqu'un test l'empêche de revenir. "
                     "Tenue à la main dans docs/suivi-tests.md.")
    feuille["A1"].font = GRAS

    colonnes = ["Feature", "Anomalie", "Trouvée par", "Conséquence si elle revenait",
                "Test qui la ferme", "État"]
    entete(feuille, colonnes, [10, 60, 24, 56, 34, 12], ligne=3)

    ligne = 4
    for anomalie in anomalies:
        for index, nom in enumerate(colonnes, start=1):
            feuille.cell(row=ligne, column=index, value=anomalie.get(nom, ""))
        ligne += 1

    plage = f"F4:F{max(ligne - 1, 4)}"
    feuille.conditional_formatting.add(plage, CellIsRule(
        operator="equal", formula=['"Close"'],
        fill=PatternFill("solid", fgColor=COULEURS_ETAT["Vert"])))
    feuille.conditional_formatting.add(plage, CellIsRule(
        operator="equal", formula=['"Ouverte"'],
        fill=PatternFill("solid", fgColor=COULEURS_ETAT["Rouge"])))
    border_tout(feuille, 3, ligne - 1, len(colonnes))


# ---------------------------------------------------------------------- main

def main() -> int:
    arguments = argparse.ArgumentParser(description=__doc__)
    arguments.add_argument("--executer", action="store_true",
                           help="relance les suites avant de générer")
    options = arguments.parse_args()

    if not SOURCE.exists():
        print(f"Source manquante : {SOURCE}", file=sys.stderr)
        return 1

    tableaux = tableaux_markdown(SOURCE.read_text(encoding="utf-8"))
    regles = [(r["Motif de chemin"].strip("`"), r["Feature"])
              for r in tableaux.get("Rattachement des suites aux features", [])]
    a_ecrire = tableaux.get("Tests à compléter", [])
    manuels = tableaux.get("Recette manuelle", [])
    anomalies = tableaux.get("Anomalies trouvées", [])

    orphelins: list[str] = []

    def rattachement(chemin: str) -> str:
        for motif, feature in regles:
            if chemin.startswith(motif):
                return feature
        orphelins.append(chemin)
        return "?"

    if options.executer:
        print("Exécution des suites…")
        executer_les_suites()

    scenarios = scenarios_backend(rattachement) + scenarios_frontend(rattachement)
    resultats, execute = lire_les_rapports()

    for scenario in scenarios:
        trouve = resultats.get(scenario.cle)
        if trouve:
            scenario.etat, scenario.duree, scenario.detail = trouve

    codes = sorted({s.feature for s in scenarios}
                   | {e.get("Feature", "") for e in a_ecrire + manuels + anomalies})
    features: dict[str, Feature] = {}
    for code in codes:
        if not code or code == "?":
            continue
        fiche = next(iter(FEATURES.glob(f"{code}-*.md")), None)
        nom, criteres = criteres_de_la_fiche(fiche) if fiche else (code, [])
        features[code] = Feature(code=code, nom=nom, criteres=criteres)

    wb = Workbook()
    feuille_synthese(wb, features, scenarios, a_ecrire, manuels, anomalies, execute)
    for feature in features.values():
        feuille_feature(wb, feature, scenarios, a_ecrire, manuels)
    feuille_couverture(wb, features, scenarios)
    feuille_anomalies(wb, anomalies)
    wb.save(SORTIE)

    verts = sum(1 for s in scenarios if s.etat == "Vert")
    rouges = sum(1 for s in scenarios if s.etat == "Rouge")
    inconnus = sum(1 for s in scenarios if s.etat == "Non exécuté")
    print(f"  {len(scenarios)} scénarios : {verts} verts, {rouges} rouges, "
          f"{inconnus} sans résultat")
    print(f"  {len(a_ecrire)} à écrire, {len(manuels)} de recette manuelle, "
          f"{len(anomalies)} anomalies")
    if orphelins:
        print("  Suites non rattachées à une feature (voir docs/suivi-tests.md) :",
              file=sys.stderr)
        for chemin in sorted(set(orphelins)):
            print("   -", chemin, file=sys.stderr)
    print(f"  écrit : {SORTIE.relative_to(RACINE)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
